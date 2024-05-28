import csv
from datetime import datetime

import openpyxl
from customtkinter import *
from categories import Categories
from CTkTable import CTkTable


def export_to_excel(filename, data):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"

    # Write data
    for row_num, row_data in enumerate(data, 1):
        for col_num, cell_value in enumerate(row_data, 1):
            sheet.cell(row=row_num, column=col_num, value=cell_value)

    workbook.save(filename)
    print(f"Data exported to {filename}")


def export_to_csv(filename, data):
    with open(filename, mode='w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        writer.writerows(data)
    print(f"Data exported to {filename}")


class ExportPage(CTkFrame):
    def __init__(self, parent, app, database, user):
        super().__init__(parent)

        self.table = None
        self.table_frame = None
        self.user_expenses = None
        self.app = app
        self.parent = parent
        self.user = user
        self.database = database

        self.add_title()
        self.add_filters()
        # self.add_export_button()
        self.get_user_expenses()
        self.add_table()
        # self.export_data()

    def add_title(self):
        title_frame = CTkFrame(master=self, fg_color="transparent")
        title_frame.pack(fill="both", padx=27, pady=(15, 0))

        CTkLabel(master=title_frame,
                 text="Export Expenses",
                 font=("Arial Black", 25),
                 text_color="#2A8C55").pack(
            anchor="nw",
            side="left")

        CTkButton(master=title_frame,
                  text="EXPORT",
                  width=150,
                  font=("Arial", 15),
                  text_color="#fff",
                  fg_color="#2A8C55",
                  hover_color="#207244",
                  command=self.export_data).pack(
            anchor="ne",
            side="right",
            pady=(10, 0),
            padx=(0, 27))

    def add_filters(self):
        self.filter_frame = CTkFrame(master=self, fg_color="transparent")
        self.filter_frame.pack(fill="both", padx=27, pady=(31, 0))

        # add start date
        CTkLabel(master=self.filter_frame,
                 text="Date:",
                 font=("Arial Bold", 17),
                 text_color="#52A476",
                 justify="left").grid(
            row=0,
            column=0,
            sticky="w")

        self.date_filter = CTkComboBox(master=self.filter_frame,
                                       width=125,
                                       values=["Date", "This month", "This year"],
                                       button_color="#2A8C55",
                                       border_color="#2A8C55",
                                       border_width=2,
                                       button_hover_color="#207244",
                                       dropdown_hover_color="#207244",
                                       dropdown_fg_color="#2A8C55",
                                       dropdown_text_color="#fff")
        self.date_filter.grid(
            row=0,
            column=1,
            sticky="w",
            padx=(5, 0))

        CTkLabel(master=self.filter_frame,
                 text="Category:",
                 font=("Arial Bold", 17),
                 text_color="#52A476",
                 justify="left").grid(
            row=0,
            column=2,
            sticky="w",
            padx=(20, 0))

        self.category_filter = CTkComboBox(master=self.filter_frame,
                                           width=125,
                                           values=[Categories.TRANSPORT.value, Categories.FOOD.value,
                                                   Categories.ENTERTAINMENT.value,
                                                   Categories.HOME.value, Categories.PERSONAL.value, ],
                                           button_color="#2A8C55",
                                           border_color="#2A8C55",
                                           border_width=2,
                                           button_hover_color="#207244",
                                           dropdown_hover_color="#207244",
                                           dropdown_fg_color="#2A8C55",
                                           dropdown_text_color="#fff")
        self.category_filter.grid(
            row=0,
            column=3,
            sticky="w",
            padx=(5, 0))

        CTkLabel(master=self.filter_frame,
                 text="Sort:",
                 font=("Arial Bold", 17),
                 text_color="#52A476",
                 justify="left").grid(
            row=0,
            column=4,
            sticky="w",
            padx=(20, 0))

        self.sort_filter = CTkComboBox(master=self.filter_frame,
                                       width=125,
                                       values=['Ascending', 'Descending'],
                                       button_color="#2A8C55",
                                       border_color="#2A8C55",
                                       border_width=2,
                                       button_hover_color="#207244",
                                       dropdown_hover_color="#207244",
                                       dropdown_fg_color="#2A8C55",
                                       dropdown_text_color="#fff")
        self.sort_filter.grid(
            row=0,
            column=5,
            sticky="w",
            padx=(5, 0))

    # def add_buttons(self):
    #     button_frame = CTkFrame(master=self, fg_color="transparent")
    #     button_frame.pack(fill="both", padx=27, pady=(15, 0))
    #
    #     CTkButton(master=button_frame,
    #               text="Apply Filters",
    #               width=150,
    #               font=("Arial Black", 15),
    #               text_color="#fff",
    #               fg_color="#2A8C55",
    #               hover_color="#207244",
    #               command=self.update_table).pack(
    #         side="left",
    #         padx=(0, 10))

    def add_table(self):
        self.table_frame = CTkScrollableFrame(master=self, fg_color="transparent")
        self.table_frame.pack(expand=True, fill="both", padx=27, pady=21)

        self.table = CTkTable(master=self.table_frame,
                              values=self.user_expenses,
                              colors=["#E6E6E6", "#EEEEEE"],
                              header_color="#2A8C55",
                              hover_color="#B4B4B4")

        self.table.edit_row(0, text_color="#fff", hover_color="#2A8C55")
        self.table.pack(expand=True)

    def get_user_expenses(self):
        self.user_expenses = []
        user_id = self.database.get_user_id(self.user.username)
        self.user_expenses = self.database.get_expenses(user_id)
        for i in range(len(self.user_expenses)):
            self.user_expenses[i] = list(self.user_expenses[i])[2:]

        self.user_expenses = [['Category', 'Amount', 'Description', 'Payment Method', 'Date']] + self.user_expenses
        print('----------------')
        print(self.user_expenses)

    def export_data(self):

        current_date = datetime.now().strftime('%Y-%m-%d')
        default_filename = f"{self.user.username}_{current_date}"

        file_path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                                 filetypes=[("Excel files", "*.xlsx"),
                                                            ("CSV files", "*.csv")],
                                                 initialfile=default_filename)

        if file_path:
            if file_path.endswith('.xlsx'):
                export_to_excel(file_path, self.user_expenses)
            elif file_path.endswith('.csv'):
                export_to_csv(file_path, self.user_expenses)

    # def update_table(self):
    #     self.get_user_expenses()
    #     filtered_expenses = self.apply_filters(self.user_expenses[1:])
    #     self.table.update_data(filtered_expenses)
    #
    # def apply_filters(self, expenses):
    #     # Placeholder method to apply filters to the expenses list
    #     # Apply date, category, and sort filters here
    #     return expenses
